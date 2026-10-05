#!/usr/bin/env python3
"""設計書(Excel)を、決定論的に JSON と Markdown に変換する。

使い方:
    python tools/design_to_json.py <xlsx> <out_root>

出力(out_root/{ブック名}/):
    index.json / index.md        ブックの概要(シート一覧、元ファイルのハッシュ、取り込めなかった情報)
    {シート名}.json              シートの正本データ
    {シート名}.md                JSON から生成した閲覧用の表(行番号 L{行}、図形番号 S{番号} 付き)
    images/                      シートに貼られた画像(図形の一覧から参照される)

取り込む情報:
    - セルの値(結合セルは範囲内の全セルに展開)
    - 数式(式そのものと、保存されている計算結果)
    - 書式の目印: 取り消し線、文字色、背景色、セル内の一部だけの取り消し線・文字色(リッチテキスト)
    - 表示形式(General 以外。例: 先頭ゼロの "000")
    - セルのコメント(メモ)、ハイパーリンク
    - 入力規則(プルダウンの選択肢など)
    - 図形・テキストボックス・グループ図形の文字と位置、コネクタ(矢印線)の位置
    - 画像(ファイルとして書き出し、位置を記録)、グラフの位置
    - 非表示のシート・行・列
    取り込めない情報(index.json の not_captured に記録):
    - 画像の中に描かれた文字・図(画像ファイルとして渡すので、必要なら目で確認する)
    - 埋め込みオブジェクト(他の Excel・Word・PDF など)の中身
    - スレッド形式のコメント(新しいコメント機能)の返信

決定論的であること:
    同じ Excel を入力すれば、何度実行しても1バイトも違わない出力になる。
    - 実行日時・実行環境に依存する値を出力しない
    - 行・列・図形などはすべて決まった順に並べる
    - JSON はキーをソートし、インデントと改行を固定する
    - 値の正規化ルール(normalize)を固定する
    このスクリプトを変更したら、SCRIPT_VERSION を上げること。
"""
import datetime as dt
import hashlib
import json
import posixpath
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

SCRIPT_VERSION = "2.0.0"

try:
    import openpyxl
    from openpyxl.utils import get_column_letter, range_boundaries
except ImportError:
    sys.exit("openpyxl がありません。pip install openpyxl を実行してください。")

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
except ImportError:  # openpyxl 3.0 系
    CellRichText, TextBlock = (), ()

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

DEFAULT_FONT_COLORS = {None, "FF000000", "theme:1", "indexed:8", "indexed:64"}
DEFAULT_FILL_COLORS = {None, "FFFFFFFF", "00000000", "theme:0", "indexed:9", "indexed:64"}


# ---------- 値と書式の正規化 ----------

def normalize(v):
    if v is None:
        return None
    if CellRichText and isinstance(v, CellRichText):
        v = str(v)
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else repr(v)
    if isinstance(v, dt.datetime):
        if v.hour == v.minute == v.second == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    s = str(v).replace("\r\n", "\n").replace("\r", "\n")
    return s if s.strip() != "" else None


def color_str(c):
    if c is None:
        return None
    t = getattr(c, "type", None)
    if t == "rgb":
        return c.rgb if isinstance(c.rgb, str) else None
    if t == "theme":
        tint = round(c.tint or 0, 4)
        return f"theme:{c.theme}" + (f"/tint:{tint}" if tint else "")
    if t == "indexed":
        return f"indexed:{c.indexed}"
    return None


def cell_attrs(cell, formula_cell):
    a = {}
    f = cell.font
    if f is not None:
        if f.strike:
            a["strike"] = True
        fc = color_str(f.color)
        if fc not in DEFAULT_FONT_COLORS:
            a["font_color"] = fc
    fill = cell.fill
    if fill is not None and fill.patternType == "solid":
        bg = color_str(fill.fgColor)
        if bg not in DEFAULT_FILL_COLORS:
            a["fill_color"] = bg
    if cell.number_format and cell.number_format != "General" and cell.value is not None \
            and not isinstance(cell.value, str):
        a["number_format"] = cell.number_format
    fv = formula_cell.value
    if isinstance(fv, str) and fv.startswith("="):
        a["formula"] = fv
    if cell.comment is not None:
        a["comment"] = normalize(cell.comment.text) or ""
    if cell.hyperlink is not None:
        a["hyperlink"] = cell.hyperlink.target or (f"#{cell.hyperlink.location}" if cell.hyperlink.location else "")
    if CellRichText and isinstance(cell.value, CellRichText):
        runs, marked = [], False
        for part in cell.value:
            if isinstance(part, TextBlock):
                run = {"text": normalize(part.text) or ""}
                ifont = part.font
                if ifont is not None and ifont.strike:
                    run["strike"] = True
                    marked = True
                rc = color_str(ifont.color) if ifont is not None else None
                if rc not in DEFAULT_FONT_COLORS:
                    run["font_color"] = rc
                    marked = True
                runs.append(run)
            else:
                runs.append({"text": normalize(str(part)) or ""})
        if marked:
            a["runs"] = runs
    return a


# ---------- 図形・画像(xlsx の XML を直接読む) ----------

def _rels(z, path):
    d = posixpath.dirname(path)
    rp = posixpath.join(d, "_rels", posixpath.basename(path) + ".rels")
    if rp not in z.namelist():
        return {}
    root = ET.fromstring(z.read(rp))
    out = {}
    for r in root.findall("rel:Relationship", NS):
        target = r.get("Target")
        if r.get("TargetMode") != "External":
            target = posixpath.normpath(posixpath.join(d, target)) if not target.startswith("/") else target.lstrip("/")
        out[r.get("Id")] = (r.get("Type", "").rsplit("/", 1)[-1], target)
    return out


def sheet_parts(z):
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = _rels(z, "xl/workbook.xml")
    out = {}
    for s in wb.findall("main:sheets/main:sheet", NS):
        rid = s.get(f"{{{NS['r']}}}id")
        if rid in rels:
            out[s.get("name")] = rels[rid][1]
    return out


def _cell_ref(marker):
    if marker is None:
        return None
    col = int(marker.find("xdr:col", NS).text)
    row = int(marker.find("xdr:row", NS).text)
    return f"{get_column_letter(col + 1)}{row + 1}"


def _text_of(el):
    paras = []
    for p in el.iter(f"{{{NS['a']}}}p"):
        paras.append("".join(t.text or "" for t in p.iter(f"{{{NS['a']}}}t")))
    s = "\n".join(paras).strip()
    return s or None


def _name_of(el):
    for tag in ("xdr:nvSpPr", "xdr:nvPicPr", "xdr:nvGrpSpPr", "xdr:nvCxnSpPr", "xdr:nvGraphicFramePr"):
        nv = el.find(tag, NS)
        if nv is not None:
            c = nv.find("xdr:cNvPr", NS)
            if c is not None:
                return c.get("name")
    return None


def _walk(el, anchor, group, drels, z, out, img_dir, img_prefix, counter):
    tag = el.tag.split("}")[-1]
    item = {"from": anchor[0], "to": anchor[1], "name": _name_of(el)}
    if group:
        item["group"] = group
    if tag == "sp":
        item["kind"] = "shape"
        item["text"] = _text_of(el)
        out.append(item)
    elif tag == "cxnSp":
        item["kind"] = "connector"
        out.append(item)
    elif tag == "pic":
        item["kind"] = "picture"
        blip = el.find(".//a:blip", NS)
        rid = blip.get(f"{{{NS['r']}}}embed") if blip is not None else None
        if rid and rid in drels and drels[rid][1] in z.namelist():
            counter[0] += 1
            ext = posixpath.splitext(drels[rid][1])[1].lower() or ".bin"
            fname = f"{img_prefix}_{counter[0]:02d}{ext}"
            img_dir.mkdir(parents=True, exist_ok=True)
            (img_dir / fname).write_bytes(z.read(drels[rid][1]))
            item["image_file"] = f"images/{fname}"
        out.append(item)
    elif tag == "graphicFrame":
        gd = el.find(".//a:graphicData", NS)
        uri = gd.get("uri", "") if gd is not None else ""
        item["kind"] = "chart" if "chart" in uri else "other"
        out.append(item)
    elif tag == "grpSp":
        gname = _name_of(el) or "group"
        for child in el:
            ctag = child.tag.split("}")[-1]
            if ctag in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
                _walk(child, anchor, (group + "/" if group else "") + gname, drels, z, out, img_dir, img_prefix, counter)


def read_drawings(z, sheet_path, img_dir, img_prefix):
    items, embedded = [], 0
    for _, (rtype, target) in sorted(_rels(z, sheet_path).items()):
        if rtype in ("oleObject", "package"):
            embedded += 1
        if rtype != "drawing" or target not in z.namelist():
            continue
        root = ET.fromstring(z.read(target))
        drels = _rels(z, target)
        counter = [0]
        for anc in root:
            atag = anc.tag.split("}")[-1]
            if atag not in ("twoCellAnchor", "oneCellAnchor", "absoluteAnchor"):
                continue
            anchor = (_cell_ref(anc.find("xdr:from", NS)), _cell_ref(anc.find("xdr:to", NS)))
            for child in anc:
                ctag = child.tag.split("}")[-1]
                if ctag in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
                    _walk(child, anchor, "", drels, z, items, img_dir, img_prefix, counter)
    for i, it in enumerate(items, 1):
        it["id"] = f"S{i}"
    return items, embedded


# ---------- シートの変換 ----------

def convert_sheet(ws_val, ws_formula, source_file, sheet_index, drawings, embedded):
    merged = sorted(str(r) for r in ws_val.merged_cells.ranges)
    merged_map = {}
    for rng in merged:
        min_col, min_row, max_col, max_row = range_boundaries(rng)
        origin = f"{get_column_letter(min_col)}{min_row}"
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                if (r, c) != (min_row, min_col):
                    merged_map[(r, c)] = (min_row, min_col, origin)

    max_row, max_col = ws_val.max_row or 0, ws_val.max_column or 0
    rows, formula_without_value, used_cols = [], [], set()
    for r in range(1, max_row + 1):
        cells, expanded, attrs = {}, {}, {}
        for c in range(1, max_col + 1):
            col = get_column_letter(c)
            if (r, c) in merged_map:
                orow, ocol, origin = merged_map[(r, c)]
                v = normalize(ws_val.cell(orow, ocol).value)
                if v is not None:
                    expanded[col] = origin
            else:
                cell = ws_val.cell(r, c)
                v = normalize(cell.value)
                a = cell_attrs(cell, ws_formula.cell(r, c))
                if a:
                    attrs[col] = a
                    used_cols.add(c)
                if v is None and "formula" in a:
                    formula_without_value.append(f"{col}{r}")
            if v is not None:
                cells[col] = v
                used_cols.add(c)
        if cells or attrs:
            row = {"row": r, "cells": cells}
            if expanded:
                row["merged_from"] = expanded
            if attrs:
                row["attrs"] = attrs
            rows.append(row)

    hidden_rows = sorted(r for r, d in ws_val.row_dimensions.items() if d.hidden)
    hidden_col_nums = set()
    for key, d in ws_val.column_dimensions.items():
        if d.hidden:
            start = d.min or range_boundaries(f"{key}1")[0]
            hidden_col_nums.update(range(start, (d.max or start) + 1))

    validations = []
    for dv in ws_val.data_validations.dataValidation:
        validations.append({"range": str(dv.sqref), "type": dv.type, "formula1": dv.formula1,
                            "formula2": dv.formula2, "allow_blank": bool(dv.allow_blank)})
    validations.sort(key=lambda x: (x["range"], x["type"] or "", x["formula1"] or ""))

    return {
        "script_version": SCRIPT_VERSION,
        "source": {"file": source_file, "sheet": ws_val.title, "sheet_index": sheet_index,
                   "sheet_state": ws_val.sheet_state},
        "columns": [get_column_letter(c) for c in sorted(used_cols)],
        "rows": rows,
        "merged_ranges": merged,
        "hidden_rows": hidden_rows,
        "hidden_columns": [get_column_letter(c) for c in sorted(hidden_col_nums)],
        "data_validations": validations,
        "drawings": drawings,
        "embedded_objects": embedded,
        "formula_without_value": formula_without_value,
    }


# ---------- Markdown ----------

def md_escape(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return str(v).replace("|", "\\|").replace("\n", "<br>")


def md_cell(v, a):
    if a and "runs" in a:
        parts = []
        for run in a["runs"]:
            t = md_escape(run["text"])
            if not t:
                continue
            if run.get("strike"):
                t = f"~~{t}~~"
            if run.get("font_color"):
                t = f"{t}‹文字色:{run['font_color']}›"
            parts.append(t)
        return "".join(parts)
    s = md_escape(v)
    if a and a.get("strike") and s:
        s = f"~~{s}~~"
    return s


def render_md(sheet: dict) -> str:
    src = sheet["source"]
    lines = [f"# {src['sheet']}", "",
             f"元ファイル: {src['file']} / シート: {src['sheet']}"
             + (" / 非表示シート" if src["sheet_state"] != "visible" else ""), "",
             "> このファイルは同名の .json から自動生成されています。行番号 L{行} は元のExcelの行番号、"
             "S{番号} は図形の番号です。~~取り消し線~~ はExcel上の取り消し線です。", ""]
    cols = sheet["columns"]
    if sheet["rows"]:
        hidden = set(sheet["hidden_rows"])
        lines.append("| 行 | " + " | ".join(cols) + " |")
        lines.append("|" + "---|" * (len(cols) + 1))
        for row in sheet["rows"]:
            label = f"L{row['row']}" + ("(非表示)" if row["row"] in hidden else "")
            merged, attrs = row.get("merged_from", {}), row.get("attrs", {})
            vals = []
            for i, c in enumerate(cols):
                v = row["cells"].get(c)
                prev = cols[i - 1] if i > 0 else None
                if c in merged and prev is not None and v == row["cells"].get(prev):
                    vals.append("")
                else:
                    vals.append(md_cell(v, attrs.get(c)))
            lines.append(f"| {label} | " + " | ".join(vals) + " |")
    else:
        lines.append("(セルのデータなし)")

    marks = {}
    notes = []
    for row in sheet["rows"]:
        for c in sorted(row.get("attrs", {}), key=lambda x: (len(x), x)):
            a, ref = row["attrs"][c], f"{c}{row['row']}"
            if a.get("strike"):
                marks.setdefault("取り消し線", []).append(ref)
            if a.get("font_color"):
                marks.setdefault(f"文字色 {a['font_color']}", []).append(ref)
            if a.get("fill_color"):
                marks.setdefault(f"背景色 {a['fill_color']}", []).append(ref)
            if "runs" in a:
                marks.setdefault("セル内の一部に取り消し線・文字色", []).append(ref)
            for key, label in (("comment", "コメント"), ("formula", "数式"),
                               ("number_format", "表示形式"), ("hyperlink", "リンク")):
                if key in a:
                    notes.append(f"| {ref} | {label} | {md_escape(a[key])} |")
    if marks:
        lines += ["", "## 書式の目印", "", "| 書式 | セル |", "|---|---|"]
        for k in sorted(marks):
            lines.append(f"| {k} | {', '.join(marks[k])} |")
    if notes:
        lines += ["", "## セルの補足(コメント・数式・表示形式・リンク)", "", "| セル | 種類 | 内容 |", "|---|---|---|"]
        lines += notes
    if sheet["drawings"]:
        kinds = {"shape": "図形", "connector": "線・矢印", "picture": "画像", "chart": "グラフ", "other": "その他"}
        lines += ["", "## 図形・画像", "", "| 番号 | 種類 | 位置 | 名前 | 文字・ファイル |", "|---|---|---|---|---|"]
        for d in sheet["drawings"]:
            pos = d["from"] or "-"
            if d.get("to"):
                pos += f"〜{d['to']}"
            content = md_escape(d.get("text")) or (d.get("image_file") or "")
            name = md_escape(d.get("name")) + (f"(グループ: {md_escape(d['group'])})" if d.get("group") else "")
            lines.append(f"| {d['id']} | {kinds.get(d['kind'], d['kind'])} | {pos} | {name} | {content} |")
    if sheet["data_validations"]:
        lines += ["", "## 入力規則", "", "| 範囲 | 種類 | 条件 |", "|---|---|---|"]
        for v in sheet["data_validations"]:
            lines.append(f"| {v['range']} | {v['type'] or ''} | {md_escape(v['formula1'])} |")
    extra = []
    if sheet["hidden_columns"]:
        extra.append(f"- 非表示の列: {', '.join(sheet['hidden_columns'])}")
    if sheet["formula_without_value"]:
        extra.append(f"- 計算結果が保存されていない数式セル: {', '.join(sheet['formula_without_value'])}")
    if sheet["embedded_objects"]:
        extra.append(f"- 埋め込みオブジェクトが {sheet['embedded_objects']} 件あります(中身は変換されていません)")
    if any(d["kind"] == "picture" for d in sheet["drawings"]):
        extra.append("- 画像の中の文字・図は変換されていません。images/ のファイルを確認してください")
    if extra:
        lines += ["", "## 注意", ""] + extra
    return "\n".join(lines) + "\n"


# ---------- 全体 ----------

def safe_name(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", name).strip() or "sheet"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def dump_json(obj, path: Path):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
                    encoding="utf-8", newline="\n")


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, out_root = Path(sys.argv[1]), Path(sys.argv[2])
    out = out_root / safe_name(src.stem)
    out.mkdir(parents=True, exist_ok=True)

    try:
        wb_val = openpyxl.load_workbook(src, data_only=True, rich_text=True)
    except TypeError:  # rich_text 未対応の openpyxl
        wb_val = openpyxl.load_workbook(src, data_only=True)
    wb_formula = openpyxl.load_workbook(src, data_only=False)
    z = zipfile.ZipFile(src)
    parts = sheet_parts(z)
    names_in_zip = set(z.namelist())

    index = {"script_version": SCRIPT_VERSION, "file": src.name, "sha256": sha256(src), "sheets": [],
             "not_captured": []}
    if any(n.startswith("xl/threadedComments/") for n in names_in_zip):
        index["not_captured"].append("スレッド形式のコメントの返信(最初のコメントはセルのコメントとして取り込み済み)")

    used = set()
    for i, ws in enumerate(wb_val.worksheets):
        name = safe_name(ws.title)
        if name in used:
            name = f"{name}_{i}"
        used.add(name)
        drawings, embedded = read_drawings(z, parts.get(ws.title, ""), out / "images", name)
        sheet = convert_sheet(ws, wb_formula[ws.title], src.name, i, drawings, embedded)
        dump_json(sheet, out / f"{name}.json")
        (out / f"{name}.md").write_text(render_md(sheet), encoding="utf-8", newline="\n")
        n_pic = sum(1 for d in drawings if d["kind"] == "picture")
        index["sheets"].append({
            "index": i, "sheet": ws.title, "file_stem": name, "sheet_state": ws.sheet_state,
            "data_rows": len(sheet["rows"]),
            "shapes_with_text": sum(1 for d in drawings if d.get("text")),
            "pictures": n_pic, "charts": sum(1 for d in drawings if d["kind"] == "chart"),
            "embedded_objects": embedded,
            "marked_cells": sum(1 for r in sheet["rows"] for a in r.get("attrs", {}).values()
                                if a.get("strike") or a.get("font_color") or a.get("fill_color") or "runs" in a),
            "comments": sum(1 for r in sheet["rows"] for a in r.get("attrs", {}).values() if "comment" in a),
            "hidden_rows": len(sheet["hidden_rows"]), "hidden_columns": len(sheet["hidden_columns"]),
            "formula_without_value": len(sheet["formula_without_value"]),
        })
        if n_pic:
            index["not_captured"].append(f"{ws.title}: 画像 {n_pic} 件の中の文字・図(images/ に書き出し済み)")
        if embedded:
            index["not_captured"].append(f"{ws.title}: 埋め込みオブジェクト {embedded} 件の中身")
        if sheet["formula_without_value"]:
            index["not_captured"].append(f"{ws.title}: 計算結果が保存されていない数式 {len(sheet['formula_without_value'])} 件の値")
    dump_json(index, out / "index.json")

    md = [f"# {src.name}", "", f"SHA-256: `{index['sha256']}` / 変換スクリプト {SCRIPT_VERSION}", "",
          "| No | シート | 状態 | データ行 | 図形の文字 | 画像 | 書式の目印 | コメント | 非表示の行/列 |",
          "|---|---|---|---|---|---|---|---|---|"]
    for s in index["sheets"]:
        md.append(f"| {s['index']} | [{s['sheet']}]({s['file_stem']}.md) | "
                  f"{'表示' if s['sheet_state'] == 'visible' else '非表示'} | {s['data_rows']} | "
                  f"{s['shapes_with_text']} | {s['pictures']} | {s['marked_cells']} | {s['comments']} | "
                  f"{s['hidden_rows']}/{s['hidden_columns']} |")
    md += ["", "## 取り込めなかった情報", ""]
    md += [f"- {x}" for x in index["not_captured"]] or ["なし"]
    (out / "index.md").write_text("\n".join(md) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps({"output_dir": str(out), "sheets": len(index["sheets"]), "sha256": index["sha256"],
                      "not_captured": len(index["not_captured"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
