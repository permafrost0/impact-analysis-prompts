#!/usr/bin/env python3
"""参照用の既存設計書から、シート名と見出しの一覧だけを取り出す(本文は出力しない)。

使い方:
    python tools/list_headings.py <xlsx> [--sheet シート名] [--grep 言葉]

    --sheet  見出しを取り出すシート(複数指定可)。省略するとすべてのシート
    --grep   この言葉を含む見出しだけを表示する(複数指定可。どれかを含めば表示)

見出しとみなす行(行の最初の値で判断):
    - 番号で始まる       例: 「4」「4.2 入力チェック」「第3章」「3-1.」
    - 括弧番号で始まる   例: 「(1) 概要」「(2)」
    - 記号で始まる       例: 「■ 入力チェック」「【処理概要】」「◆」

出力するのは見出しの1行目(最大80文字)と行番号だけで、本文・表・図形の文字は出力しない。
このスクリプトは Excel を読むだけで、保存・変更はしない。
"""
import argparse
import re
import sys

# Windows のコンソール(cp932)で表せない文字があっても止まらないようにする
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl がありません。pip install openpyxl を実行してください。")

HEADING = re.compile(
    r"^(?:"
    r"第\d+[章節項]"                                   # 第3章
    r"|\d+(?:[.．\-－]\d+)+[.．]?(?:[\s　]|$|(?=\D))"  # 4.2 / 4.2入力チェック / 3-1.
    r"|\d+[.．、)）]?(?:[\s　]|$)"                      # 4 会員登録 / 4. / 4)
    r"|\d+[.．、)）](?=\S)"                             # 4.会員登録
    r"|[（(]\d+[)）]"                                   # (1)
    r"|[■□◆◇●○▼▽★☆【\[]"                           # ■ 【
    r")"
)


def first_text(row):
    for v in row:
        if v is None:
            continue
        s = str(v).strip()
        if s:
            return s
    return None


def main():
    ap = argparse.ArgumentParser(description="見出しの一覧だけを取り出す(本文は出力しない)")
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", action="append", default=[])
    ap.add_argument("--grep", action="append", default=[])
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    names = a.sheet or wb.sheetnames
    missing = [s for s in names if s not in wb.sheetnames]
    if missing:
        print("見つからないシートがあります: " + ", ".join(missing), file=sys.stderr)
        print("このブックのシート: " + ", ".join(wb.sheetnames), file=sys.stderr)
        sys.exit(2)

    print(f"# {a.xlsx} の見出し(本文は含みません)")
    print()
    print("シート: " + " / ".join(wb.sheetnames))
    for name in names:
        lines = []
        for r, row in enumerate(wb[name].iter_rows(values_only=True), 1):
            s = first_text(row)
            if not s or not HEADING.match(s):
                continue
            s = s.splitlines()[0][:80]
            if a.grep and not any(g in s for g in a.grep):
                continue
            lines.append(f"L{r:<5} {s}")
        print()
        print(f"## シート: {name}")
        print()
        print("\n".join(lines) if lines else "(見出しなし)")


if __name__ == "__main__":
    main()
