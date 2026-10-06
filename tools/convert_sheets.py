#!/usr/bin/env python3
"""指示されたシートだけを、AI が読みやすい行単位の Markdown に変換する。

使い方:
    シート名の一覧だけを表示する(中身は変換しない):
        python tools/convert_sheets.py --list <xlsx>
    指示されたシートを変換する:
        python tools/convert_sheets.py <xlsx> <出力フォルダ> --sheet 画面項目定義 --sheet イベント一覧
    すべてのシートを変換する(ユーザーが「全シートが対象」と指示したときだけ):
        python tools/convert_sheets.py <xlsx> <出力フォルダ> --all

出力(<出力フォルダ>/{ブック名}/):
    {シート名}.md   シートの中身(値のあるセルだけを、行ごとに1行で)
    index.md        ブック内の全シート名と、変換したかどうか、元ファイルの SHA-256

書式:
    L3     [A]2 [B]電話番号 [C]tel   〈文字色 赤(FF0000): A〜C〉〈コメント C: 必須に変更〉
    - L3 は元の Excel の行番号。[A] は列。結合セルは [B〜D] のように範囲で1回だけ書く
    - 取り消し線は ~~ ~~、セル内の一部の色は {赤(FF0000):11} のように書く(色の名前は目安で、色コードを併記)
    - セル内の改行は ⏎
    - 図形・テキストボックスの文字は「図形」の節に、近くの行番号つきで書く

このスクリプトは Excel を読むだけで、保存・変更はしない。
同じ入力からは常に同じ出力になる(実行日時などは出力しない)。
"""
import argparse
import datetime as dt
import hashlib
import posixpath
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

# Windows のコンソール(cp932)で表せない文字があっても止まらないようにする
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass

SCRIPT_VERSION = "3.0.0"

try:
    import openpyxl
    from openpyxl.utils import get_column_letter, range_boundaries
except ImportError:
    sys.exit("openpyxl がありません。pip install openpyxl を実行してください。")

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
except ImportError:  # 古い openpyxl
    CellRichText, TextBlock = (), ()

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

# 色の名前(Excel の標準色)。名前は目安で、必ず色コードを併記する
PALETTE = [
    ("濃い赤", (0xC0, 0x00, 0x00)), ("赤", (0xFF, 0x00, 0x00)), ("オレンジ", (0xFF, 0xC0, 0x00)),
    ("黄", (0xFF, 0xFF, 0x00)), ("薄い緑", (0x92, 0xD0, 0x50)), ("緑", (0x00, 0xB0, 0x50)),
    ("薄い青", (0x00, 0xB0, 0xF0)), ("青", (0x00, 0x70, 0xC0)), ("濃い青", (0x00, 0x20, 0x60)),
    ("紫", (0x70, 0x30, 0xA0)), ("灰", (0x80, 0x80, 0x80)), ("薄い灰", (0xD9, 0xD9, 0xD9)),
    ("黒", (0x00, 0x00, 0x00)), ("白", (0xFF, 0xFF, 0xFF)),
]
DEFAULT_FONT = {None, "FF000000", "theme:1", "indexed:8", "indexed:64"}
DEFAULT_FILL = {None, "FFFFFFFF", "00000000", "theme:0", "indexed:9", "indexed:64"}


# ---------------- 値と色 ----------------

def color_key(c):
    if c is None:
        return None
    t = getattr(c, "type", None)
    if t == "rgb":
        return c.rgb if isinstance(c.rgb, str) else None
    if t == "theme":
        tint = round(c.tint or 0, 2)
        return f"theme:{c.theme}" + (f"/tint:{tint}" if tint else "")
    if t == "indexed":
        return f"indexed:{c.indexed}"
    return None


def color_name(key):
    if key is None:
        return ""
    if key.startswith("theme:"):
        return "テーマ色" + key.split(":", 1)[1]
    if key.startswith("indexed:"):
        return "色番号" + key.split(":", 1)[1]
    hexv = key[-6:].upper()
    try:
        rgb = tuple(int(hexv[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return key
    name = min(PALETTE, key=lambda p: sum((a - b) ** 2 for a, b in zip(p[1], rgb)))[0]
    return f"{name}({hexv})"


def norm_value(v):
    if v is None:
        return None
    if CellRichText and isinstance(v, CellRichText):
        v = str(v)
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, dt.datetime):
        if v.hour == v.minute == v.second == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    s = str(v).replace("\r\n", "\n").replace("\r", "\n")
    return s if s.strip() else None


def one_line(s):
    return "" if s is None else s.replace("\n", " ⏎ ")


def apply_number_format(v, fmt):
    """先頭ゼロ(000)と小数桁(0.00)だけ見た目どおりにする。それ以外は値のまま。"""
    if fmt is None or not isinstance(v, (int, float)) or isinstance(v, bool):
        return None
    if re.fullmatch(r"0+", fmt) and float(v).is_integer():
        return str(int(v)).zfill(len(fmt))
    m = re.fullmatch(r"0\.(0+)", fmt)
    if m:
        return f"{v:.{len(m.group(1))}f}"
    return None


def col_range(cols):
    """['A','B','C','E'] -> 'A〜C, E'"""
    nums = sorted({_col_num(c) for c in cols})
    parts, start, prev = [], None, None
    for n in nums:
        if start is None:
            start = prev = n
        elif n == prev + 1:
            prev = n
        else:
            parts.append((start, prev))
            start = prev = n
    if start is not None:
        parts.append((start, prev))
    return ", ".join(get_column_letter(a) if a == b else f"{get_column_letter(a)}〜{get_column_letter(b)}"
                     for a, b in parts)


def _col_num(letter):
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - 64)
    return n


# ---------------- 図形(xlsx の XML を直接読む) ----------------

def _rels(z, path):
    d = posixpath.dirname(path)
    rp = posixpath.join(d, "_rels", posixpath.basename(path) + ".rels")
    if rp not in z.namelist():
        return {}
    out = {}
    for r in ET.fromstring(z.read(rp)).findall("rel:Relationship", NS):
        target = r.get("Target")
        if r.get("TargetMode") != "External":
            target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(d, target))
        out[r.get("Id")] = (r.get("Type", "").rsplit("/", 1)[-1], target)
    return out


def sheet_parts(z):
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = _rels(z, "xl/workbook.xml")
    out = {}
    for s in wb.findall("main:sheets/main:sheet", NS):
        rid = s.get(f"{{{NS['r']}}}id")
        if rid in rels:
            out[s.get("name")] = rels[rid][1]
    return out


def _marker(m):
    if m is None:
        return None
    return int(m.find("xdr:row", NS).text) + 1, int(m.find("xdr:col", NS).text) + 1


def _text_of(el):
    paras = ["".join(t.text or "" for t in p.iter(f"{{{NS['a']}}}t")) for p in el.iter(f"{{{NS['a']}}}p")]
    s = "\n".join(paras).strip()
    return s or None


def _name_of(el):
    for tag in ("xdr:nvSpPr", "xdr:nvPicPr", "xdr:nvGrpSpPr", "xdr:nvCxnSpPr", "xdr:nvGraphicFramePr"):
        nv = el.find(tag, NS)
        if nv is not None and nv.find("xdr:cNvPr", NS) is not None:
            return nv.find("xdr:cNvPr", NS).get("name")
    return None


def _walk(el, frm, to, group, out):
    tag = el.tag.split("}")[-1]
    item = {"from": frm, "to": to, "name": _name_of(el), "group": group}
    if tag == "sp":
        item["kind"], item["text"] = "shape", _text_of(el)
        out.append(item)
    elif tag == "cxnSp":
        item["kind"] = "connector"
        out.append(item)
    elif tag == "pic":
        item["kind"] = "picture"
        out.append(item)
    elif tag == "graphicFrame":
        gd = el.find(".//a:graphicData", NS)
        item["kind"] = "chart" if gd is not None and "chart" in gd.get("uri", "") else "other"
        out.append(item)
    elif tag == "grpSp":
        gname = _name_of(el) or "グループ"
        for child in el:
            if child.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
                _walk(child, frm, to, (group + "/" if group else "") + gname, out)


def read_drawings(z, sheet_path):
    items, embedded = [], 0
    for _, (rtype, target) in sorted(_rels(z, sheet_path).items()):
        if rtype in ("oleObject", "package"):
            embedded += 1
        if rtype != "drawing" or target not in z.namelist():
            continue
        for anc in ET.fromstring(z.read(target)):
            if anc.tag.split("}")[-1] not in ("twoCellAnchor", "oneCellAnchor", "absoluteAnchor"):
                continue
            frm, to = _marker(anc.find("xdr:from", NS)), _marker(anc.find("xdr:to", NS))
            for child in anc:
                if child.tag.split("}")[-1] in ("sp", "cxnSp", "pic", "graphicFrame", "grpSp"):
                    _walk(child, frm, to, "", items)
    for i, it in enumerate(items, 1):
        it["id"] = f"S{i}"
    return items, embedded


def _ref(pos):
    return f"{get_column_letter(pos[1])}{pos[0]}" if pos else "-"


# ---------------- シートの変換 ----------------

def convert_sheet(ws, ws_f, book_name, drawings, embedded):
    # 結合セル: 左上のセル → (最終行, 最終列)
    merged = {}
    for rng in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(rng))
        merged[(r1, c1)] = (r2, c2)

    hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
    hidden_cols = set()
    for key, d in ws.column_dimensions.items():
        if d.hidden:
            start = d.min or range_boundaries(f"{key}1")[0]
            hidden_cols.update(range(start, (d.max or start) + 1))

    rows = {}
    counts = {"marks": 0, "comments": 0}
    # 実在するセルだけを調べる(書式だけが広い範囲に付いていても遅くならない)
    for (r, c) in sorted(ws._cells.keys()):
        cell = ws._cells[(r, c)]
        if type(cell).__name__ == "MergedCell":
            continue
        raw = cell.value
        fcell = ws_f._cells.get((r, c))
        formula = fcell.value if fcell is not None and isinstance(fcell.value, str) and fcell.value.startswith("=") else None
        comment = norm_value(cell.comment.text) if cell.comment is not None else None
        link = None
        if cell.hyperlink is not None:
            link = cell.hyperlink.target or (f"#{cell.hyperlink.location}" if cell.hyperlink.location else None)
        font = cell.font
        strike = bool(font is not None and font.strike)
        fcolor = color_key(font.color) if font is not None else None
        fcolor = None if fcolor in DEFAULT_FONT else fcolor
        fill = None
        if cell.fill is not None and cell.fill.patternType == "solid":
            fill = color_key(cell.fill.fgColor)
            fill = None if fill in DEFAULT_FILL else fill

        # 値の文字列(セル内の一部の書式を反映)
        text = None
        partial = False
        if CellRichText and isinstance(raw, CellRichText):
            parts = []
            for part in raw:
                if isinstance(part, TextBlock):
                    t = one_line(norm_value(part.text) or "")
                    if not t:
                        continue
                    f = part.font
                    if f is not None and f.strike:
                        t, partial = f"~~{t}~~", True
                    ck = color_key(f.color) if f is not None else None
                    if ck not in DEFAULT_FONT and ck is not None:
                        t, partial = f"{{{color_name(ck)}:{t}}}", True
                    parts.append(t)
                else:
                    parts.append(one_line(norm_value(str(part)) or ""))
            text = "".join(parts) or None
        else:
            v = norm_value(raw)
            shown = apply_number_format(raw, cell.number_format) if v is not None else None
            text = one_line(shown or v) if v is not None else None
        if formula is not None:
            text = f"{text}(数式 {formula})" if text is not None else f"(数式 {formula} / 計算結果なし)"

        has_mark = strike or fcolor or fill or partial
        if text is None and not (has_mark or comment or link):
            continue
        end = merged.get((r, c))
        label = get_column_letter(c)
        span_cols = [get_column_letter(x) for x in range(c, (end[1] if end else c) + 1)]
        if end and end[1] > c:
            label = f"{get_column_letter(c)}〜{get_column_letter(end[1])}"
        if end and end[0] > r:
            label += f"(〜L{end[0]})"
        if c in hidden_cols:
            label += "(非表示)"
        if strike and text and not text.startswith("~~"):
            text = f"~~{text}~~"
        rows.setdefault(r, []).append({
            "label": label, "text": text or "", "cols": span_cols, "strike": strike, "fcolor": fcolor,
            "fill": fill, "partial": partial, "comment": comment, "link": link,
        })
        counts["marks"] += 1 if has_mark else 0
        counts["comments"] += 1 if comment else 0

    # 本文
    body = []
    for r in sorted(rows):
        cells = rows[r]
        head = f"L{r}" + ("(非表示)" if r in hidden_rows else "")
        line = f"{head:<6} " + " ".join(f"[{x['label']}]{x['text']}" for x in cells if x["text"])
        marks = []
        strike_cols = [c for x in cells if x["strike"] for c in x["cols"]]
        if strike_cols:
            marks.append(f"〈取り消し線: {col_range(strike_cols)}〉")
        for key in sorted({x["fcolor"] for x in cells if x["fcolor"]}):
            cols = [c for x in cells if x["fcolor"] == key for c in x["cols"]]
            marks.append(f"〈文字色 {color_name(key)}: {col_range(cols)}〉")
        for key in sorted({x["fill"] for x in cells if x["fill"]}):
            cols = [c for x in cells if x["fill"] == key for c in x["cols"]]
            marks.append(f"〈背景 {color_name(key)}: {col_range(cols)}〉")
        for x in cells:
            if x["comment"]:
                marks.append(f"〈コメント {x['cols'][0]}: {one_line(x['comment'])}〉")
            if x["link"]:
                marks.append(f"〈リンク {x['cols'][0]}: {x['link']}〉")
        body.append(line.rstrip() + ("   " + "".join(marks) if marks else ""))

    # 図形
    shapes = [d for d in drawings if d["kind"] == "shape" and d.get("text")]
    shape_lines = []
    for d in shapes:
        pos = _ref(d["from"]) + (f"〜{_ref(d['to'])}" if d.get("to") else "")
        near = f"L{d['from'][0]}" + (f"〜L{d['to'][0]}" if d.get("to") and d["to"][0] != d["from"][0] else "") \
            if d.get("from") else "位置不明"
        grp = f"(グループ「{d['group']}」)" if d.get("group") else ""
        shape_lines.append(f"{d['id']:<4} {pos}({near} の位置){grp}: {one_line(d['text'])}")

    # 入力規則
    dv_lines = []
    for dv in sorted(ws.data_validations.dataValidation, key=lambda x: (str(x.sqref), x.type or "")):
        rng = str(dv.sqref).replace(":", "〜")
        if dv.type == "list":
            dv_lines.append(f"{rng}: 選択肢 {dv.formula1}")
        else:
            cond = " ".join(x for x in (dv.operator, dv.formula1, dv.formula2) if x)
            dv_lines.append(f"{rng}: {dv.type or '指定なし'} {cond}".rstrip())

    notes = []
    if hidden_cols:
        notes.append(f"非表示の列: {col_range([get_column_letter(c) for c in hidden_cols])}")
    pics = [d for d in drawings if d["kind"] == "picture"]
    if pics:
        notes.append(f"画像 {len(pics)} 件({', '.join(_ref(d['from']) for d in pics)})の中身は読み取っていません")
    charts = [d for d in drawings if d["kind"] == "chart"]
    if charts:
        notes.append(f"グラフ {len(charts)} 件の中身は読み取っていません")
    conns = [d for d in drawings if d["kind"] == "connector"]
    if conns:
        notes.append(f"線・矢印 {len(conns)} 件(図形どうしのつながりは読み取っていません)")
    if embedded:
        notes.append(f"埋め込みオブジェクト {embedded} 件の中身は読み取っていません")

    used = [c for r in rows for x in rows[r] for c in x["cols"]]
    summary = (f"範囲: L{min(rows)}〜L{max(rows)} / 列 {col_range(used)}" if rows else "範囲: (値のあるセルなし)")
    summary += f" / 図形 {len(shapes)} / 書式の目印 {counts['marks']} / コメント {counts['comments']}"
    if ws.sheet_state != "visible":
        summary += " / 非表示のシート"

    out = [f"# {ws.title}({book_name})", "", summary, "", "## セル", ""]
    out += body or ["(値のあるセルはありません)"]
    if shape_lines:
        out += ["", "## 図形・テキストボックス", ""] + shape_lines
    if dv_lines:
        out += ["", "## 入力規則", ""] + dv_lines
    if notes:
        out += ["", "## 注意", ""] + [f"- {n}" for n in notes]
    return "\n".join(out) + "\n"


def safe_name(name):
    return re.sub(r'[\\/:*?"<>|]', "_", name).strip() or "sheet"


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def main():
    ap = argparse.ArgumentParser(description="指示されたシートだけを Markdown に変換する")
    ap.add_argument("--list", metavar="XLSX", help="シート名の一覧だけを表示する")
    ap.add_argument("xlsx", nargs="?")
    ap.add_argument("out_dir", nargs="?")
    ap.add_argument("--sheet", action="append", default=[], help="変換するシート名(複数指定可)")
    ap.add_argument("--all", action="store_true", help="すべてのシートを変換する")
    a = ap.parse_args()

    if a.list:
        wb = openpyxl.load_workbook(a.list, read_only=True)
        for i, name in enumerate(wb.sheetnames, 1):
            print(f"{i}\t{name}")
        return

    if not a.xlsx or not a.out_dir or (not a.sheet and not a.all):
        ap.error("<xlsx> <出力フォルダ> と、--sheet または --all が必要です")

    src = Path(a.xlsx)
    try:
        wb = openpyxl.load_workbook(src, data_only=True, rich_text=True)
    except TypeError:
        wb = openpyxl.load_workbook(src, data_only=True)
    names = wb.sheetnames
    targets = names if a.all else a.sheet
    missing = [s for s in targets if s not in names]
    if missing:
        print("見つからないシートがあります: " + ", ".join(missing), file=sys.stderr)
        print("このブックのシート: " + ", ".join(names), file=sys.stderr)
        sys.exit(2)

    wb_f = openpyxl.load_workbook(src, data_only=False)
    z = zipfile.ZipFile(src)
    parts = sheet_parts(z)
    out = Path(a.out_dir) / safe_name(src.stem)
    out.mkdir(parents=True, exist_ok=True)

    for name in targets:
        drawings, embedded = read_drawings(z, parts.get(name, ""))
        md = convert_sheet(wb[name], wb_f[name], src.name, drawings, embedded)
        (out / f"{safe_name(name)}.md").write_text(md, encoding="utf-8", newline="\n")
        print(f"変換: {out / (safe_name(name) + '.md')}")

    lines = [f"# {src.name}", "", f"SHA-256: `{sha256(src)}` / 変換スクリプト {SCRIPT_VERSION}", "",
             "| No | シート | 変換 |", "|---|---|---|"]
    for i, name in enumerate(names, 1):
        done = (out / f"{safe_name(name)}.md").exists()
        lines.append(f"| {i} | {name} | {'[変換済み](' + safe_name(name) + '.md)' if done else '-'} |")
    lines += ["", "変換するのは指示されたシートだけです。「-」のシートは読んでいません。"]
    (out / "index.md").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
